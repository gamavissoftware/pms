from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont
import textwrap
from pathlib import Path


BASE_DIR = Path(__file__).resolve().parent
DOCX_PATH = BASE_DIR / "SF_Module_Flow_Review_for_PMS.docx"
DIAGRAM_PATH = BASE_DIR / "sf_module_flow_diagram.png"


FONT_CANDIDATES = [
    "/System/Library/Fonts/Supplemental/Arial.ttf",
    "/System/Library/Fonts/Supplemental/Helvetica.ttf",
    "/Library/Fonts/Arial.ttf",
]
BOLD_FONT_CANDIDATES = [
    "/System/Library/Fonts/Supplemental/Arial Bold.ttf",
    "/System/Library/Fonts/Supplemental/Helvetica Bold.ttf",
    "/Library/Fonts/Arial Bold.ttf",
]


def load_font(size, bold=False):
    candidates = BOLD_FONT_CANDIDATES if bold else FONT_CANDIDATES
    for candidate in candidates:
        if Path(candidate).exists():
            return ImageFont.truetype(candidate, size=size)
    return ImageFont.load_default()


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell, color="D9D9D9", size="6"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_text(cell, text, bold=False, size=9, color="000000"):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Arial"
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def remove_paragraph_borders_from_style(style):
    p_pr = style.element.get_or_add_pPr()
    for child in list(p_pr):
        if child.tag == qn("w:pBdr"):
            p_pr.remove(child)


def disable_paragraph_borders(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    for child in list(p_pr):
        if child.tag == qn("w:pBdr"):
            p_pr.remove(child)
    p_bdr = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "nil")
        p_bdr.append(element)
    p_pr.append(p_bdr)


def style_table(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = widths[index]
            set_cell_border(cell)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.08


def add_bullets(doc, items):
    for item in items:
        para = doc.add_paragraph(style="List Bullet")
        para.paragraph_format.space_after = Pt(3)
        run = para.add_run(item)
        run.font.name = "Arial"
        run.font.size = Pt(10.5)


def draw_wrapped_text(draw, xy, text, font, fill, max_width, align="center", line_gap=7):
    words = text.split()
    lines = []
    line = ""
    for word in words:
        probe = word if not line else f"{line} {word}"
        if draw.textbbox((0, 0), probe, font=font)[2] <= max_width:
            line = probe
        else:
            if line:
                lines.append(line)
            line = word
    if line:
        lines.append(line)

    line_heights = []
    total_height = 0
    for line in lines:
        bbox = draw.textbbox((0, 0), line, font=font)
        h = bbox[3] - bbox[1]
        line_heights.append(h)
        total_height += h
    total_height += line_gap * max(0, len(lines) - 1)

    x, y, width, height = xy
    current_y = y + (height - total_height) / 2
    for line, h in zip(lines, line_heights):
        bbox = draw.textbbox((0, 0), line, font=font)
        text_width = bbox[2] - bbox[0]
        if align == "left":
            text_x = x
        elif align == "right":
            text_x = x + width - text_width
        else:
            text_x = x + (width - text_width) / 2
        draw.text((text_x, current_y), line, fill=fill, font=font)
        current_y += h + line_gap


def box(draw, xy, title, body="", fill="#FFFFFF", outline="#8EAADB"):
    x, y, w, h = xy
    draw.rounded_rectangle((x, y, x + w, y + h), radius=18, fill=fill, outline=outline, width=3)
    title_font = load_font(28, bold=True)
    body_font = load_font(22)
    if body:
        draw_wrapped_text(draw, (x + 24, y + 10, w - 48, 36), title, title_font, "#111827", w - 48)
        draw_wrapped_text(draw, (x + 24, y + 44, w - 48, h - 52), body, body_font, "#374151", w - 48, line_gap=5)
    else:
        draw_wrapped_text(draw, (x + 24, y + 12, w - 48, h - 24), title, title_font, "#111827", w - 48)


def arrow(draw, start, end, color="#4B5563", width=4):
    draw.line((start, end), fill=color, width=width)
    ex, ey = end
    sx, sy = start
    if abs(ey - sy) >= abs(ex - sx):
        direction = 1 if ey >= sy else -1
        points = [(ex, ey), (ex - 12, ey - direction * 20), (ex + 12, ey - direction * 20)]
    else:
        direction = 1 if ex >= sx else -1
        points = [(ex, ey), (ex - direction * 20, ey - 12), (ex - direction * 20, ey + 12)]
    draw.polygon(points, fill=color)


def draw_diagram():
    width, height = 1700, 1420
    img = Image.new("RGB", (width, height), "#FFFFFF")
    draw = ImageDraw.Draw(img)
    title_font = load_font(42, bold=True)
    small_font = load_font(24)

    draw.text((70, 50), "SF Module Flow From Quotation to Dispatch", fill="#111827", font=title_font)
    draw.text((72, 108), "Review draft for PMS team approval before development", fill="#4B5563", font=small_font)

    left_x = 95
    right_x = 955
    box_w = 650
    box_h = 120
    gap = 18
    top_y = 170

    steps = [
        ("Quotation", "Select quotation type: Consumable, Critical, Cons + Critical, or Custom Engg. Select Self Pickup or Courier."),
        ("Custom Engg Check", "If Custom Engg, capture subtype: Changeover or Speed Upgradation."),
        ("Order Own", "Assign internal ownership and convert approved quotation to order record."),
        ("Upload PO", "Customer PO is mandatory before SF creation."),
        ("Create SF", "Tripti owner. TAT 1 day. Capture country wise document requirements and packing and dispatch duration."),
        ("Order Dashboard", "Run MRP on same day as SF creation. Owner Jitender."),
        ("Shortage Report", "Generate same day once SF is created. Owner Jitender."),
        ("Design TAT", "Complete BOM within 1 week where design/BOM is required."),
        ("Packing and Logistics", "Calculate packing boxes and arrange logistics for domestic or international dispatch."),
        ("Raise Indent", "Raise indent to Purchase same day. Owner Jitender."),
        ("Purchase PO", "Purchase raises Mechanical PO, BOP PO, Sec 6 PO, and Automation PO."),
        ("Store Receipt", "Store receives items group wise. Target TAT 3 weeks."),
        ("CE Assembly", "For Custom Engg only, assembly TAT 3 to 4 days."),
        ("Dispatch", "Check uploaded dispatch documents before release."),
        ("Acknowledgement", "Record acknowledgement of goods received."),
    ]

    centers = {}
    for idx, (title, body) in enumerate(steps):
        col_x = left_x if idx < 8 else right_x
        row = idx if idx < 8 else idx - 8
        y = top_y + row * (box_h + gap)
        fill = "#EAF2FF"
        outline = "#7EA6D8"
        if idx in (1, 12):
            fill = "#FFF4E5"
            outline = "#D9A441"
        if idx in (10, 11):
            fill = "#EEF8F1"
            outline = "#76B985"
        if idx in (13, 14):
            fill = "#F2F2F2"
            outline = "#9CA3AF"
        box(draw, (col_x, y, box_w, box_h), title, body, fill, outline)
        centers[idx] = (col_x + box_w / 2, y + box_h / 2)

    for i in range(0, 7):
        arrow(draw, (centers[i][0], centers[i][1] + box_h / 2 - 4), (centers[i + 1][0], centers[i + 1][1] - box_h / 2 + 4))
    arrow(draw, (centers[7][0] + box_w / 2, centers[7][1]), (centers[8][0] - box_w / 2 + 6, centers[8][1]))
    for i in range(8, len(centers) - 1):
        arrow(draw, (centers[i][0], centers[i][1] + box_h / 2 - 4), (centers[i + 1][0], centers[i + 1][1] - box_h / 2 + 4))

    img.save(DIAGRAM_PATH, quality=95)


def setup_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    styles = doc.styles
    for style_name in ("Normal", "Title", "Heading 1", "Heading 2"):
        style = styles[style_name]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor(0, 0, 0)
        remove_paragraph_borders_from_style(style)
    styles["Normal"].font.size = Pt(10.5)
    styles["Title"].font.size = Pt(24)
    styles["Heading 1"].font.size = Pt(15)
    styles["Heading 2"].font.size = Pt(12)
    return doc


def add_header(doc):
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    disable_paragraph_borders(title)
    title.add_run("SF Module Flow Review for PMS")
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(10)
    run = subtitle.add_run("From spares quotation to order dispatch and goods acknowledgement")
    run.font.name = "Arial"
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(80, 80, 80)


def add_scope(doc):
    doc.add_heading("Purpose", level=1)
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(8)
    run = para.add_run(
        "This document captures the proposed SF process for the PMS spares module so the PMS team can review and approve the business flow before development starts. "
        "The requested flow replaces the earlier built spares execution flow for this module after review approval."
    )
    run.font.name = "Arial"
    run.font.size = Pt(10.5)

    doc.add_heading("Scope Of This Flow", level=1)
    add_bullets(
        doc,
        [
            "Starts from quotation finalization in the spares module.",
            "Covers quotation type, customer PO upload, SF creation, MRP, shortage, design BOM, purchase, store receipt, assembly where applicable, dispatch, and customer acknowledgement.",
            "Applies to domestic and international dispatch, with document requirements captured country wise.",
            "Keeps Custom Engg cases separate where Changeover or Speed Upgradation is selected.",
        ],
    )

    doc.add_heading("Mandatory Selection Fields", level=1)
    table = doc.add_table(rows=1, cols=3)
    widths = [Inches(1.65), Inches(3.35), Inches(2.1)]
    style_table(table, widths)
    headers = ["Field", "Allowed Values", "Usage"]
    for idx, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[idx], header, bold=True, size=9.5, color="FFFFFF")
        set_cell_shading(table.rows[0].cells[idx], "1F4E79")
    rows = [
        ("Quotation Type", "Consumable, Critical, Cons + Critical, Custom Engg.", "Required at quotation stage."),
        ("Dispatch Mode", "Self Pickup, Courier", "Required at quotation stage."),
        ("Custom Engg Type", "Changeover, Speed Upgradation", "Required only when quotation type is Custom Engg."),
    ]
    for row_index, row in enumerate(rows, start=1):
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], value, size=9)
            if row_index % 2 == 0:
                set_cell_shading(cells[idx], "F3F7FB")


def add_diagram_page(doc):
    doc.add_page_break()
    doc.add_heading("Process Flow Diagram", level=1)
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(6)
    run = para.add_run(
        "The diagram below shows the proposed operational sequence and where conditional Custom Engg handling enters the flow."
    )
    run.font.name = "Arial"
    run.font.size = Pt(10.5)
    pic = doc.add_picture(str(DIAGRAM_PATH), width=Inches(6.85))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_stage_matrix(doc):
    doc.add_page_break()
    doc.add_heading("Stage Matrix For Review", level=1)
    intro = doc.add_paragraph()
    intro.paragraph_format.space_after = Pt(8)
    run = intro.add_run("This matrix converts the flow into PMS implementation checkpoints for review.")
    run.font.name = "Arial"
    run.font.size = Pt(10.5)

    rows = [
        ("1", "Quotation setup", "Select quotation type, dispatch mode, and Custom Engg subtype if applicable.", "Sales or spares user", "Before order conversion", "Add mandatory fields and validation."),
        ("2", "Order ownership", "Assign internal order owner and confirm customer PO.", "Order owner", "Before SF creation", "Create order ownership field and PO upload gate."),
        ("3", "Create SF", "Create SF with document requirements country wise and packing and dispatch duration.", "Tripti", "1 day", "SF form must capture domestic or international handling."),
        ("4", "Run MRP", "Run MRP from order dashboard after SF is created.", "Jitender", "Same day", "Add dashboard action and date stamp."),
        ("5", "Shortage report", "Generate shortage report once SF is created.", "Jitender", "Same day", "Link shortage report to SF and MRP result."),
        ("6", "Design and BOM", "Complete design requirement and BOM, mainly for Custom Engg where needed.", "Design team", "1 week", "Track BOM completion and design TAT."),
        ("7", "Packing and logistics", "Calculate packing boxes and arrange logistics.", "Planning or logistics", "Before dispatch readiness", "Capture box count and logistics status."),
        ("8", "Raise indent", "Raise indent to Purchase.", "Jitender", "Same day", "Indent should be generated from approved requirement."),
        ("9", "Purchase PO", "Raise Mechanical PO, BOP PO, Sec 6 PO, and Automation PO.", "Purchase department", "As per purchase process", "Track each PO group separately."),
        ("10", "Store receipt", "Receive items group wise.", "Store", "3 weeks", "Store receipt must update group wise readiness."),
        ("11", "Assembly for CE", "Complete assembly for Custom Engg cases.", "Assembly team", "3 to 4 days", "Conditional stage when Custom Engg is selected."),
        ("12", "Dispatch", "Verify uploaded dispatch documents before dispatch.", "Dispatch team", "Before release", "Document checklist must be mandatory."),
        ("13", "Acknowledgement", "Capture acknowledgement that goods were received.", "Customer or PMS user", "After delivery", "Add goods received acknowledgement upload or confirmation."),
    ]
    table = doc.add_table(rows=1, cols=6)
    widths = [Inches(0.42), Inches(1.3), Inches(2.15), Inches(1.0), Inches(0.9), Inches(1.55)]
    style_table(table, widths)
    headers = ["No", "Stage", "Action", "Owner", "TAT", "PMS Requirement"]
    for idx, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[idx], header, bold=True, size=8.5, color="FFFFFF")
        set_cell_shading(table.rows[0].cells[idx], "1F4E79")

    for row_index, row in enumerate(rows, start=1):
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], value, size=7.8)
            if idx in (0, 4):
                cells[idx].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            if row_index % 2 == 0:
                set_cell_shading(cells[idx], "F3F7FB")


def add_review_points(doc):
    doc.add_heading("Review Points Before Development", level=1)
    add_bullets(
        doc,
        [
            "Confirm meaning and responsible person for Order Own.",
            "Confirm country wise document checklist fields for domestic and international dispatch.",
            "Confirm whether MRP and shortage report are system generated, manually uploaded, or both.",
            "Confirm if design TAT applies to all quotation types or only Custom Engg cases.",
            "Confirm purchase PO groups and whether Sec 6 should be shown as a separate vendor or internal department.",
            "Confirm acknowledgement format: checkbox, uploaded proof, received date, receiver name, or all of these.",
        ],
    )


def main():
    draw_diagram()
    doc = setup_document()
    add_header(doc)
    add_scope(doc)
    add_diagram_page(doc)
    add_stage_matrix(doc)
    add_review_points(doc)
    doc.save(DOCX_PATH)
    print(DOCX_PATH)
    print(DIAGRAM_PATH)


if __name__ == "__main__":
    main()
